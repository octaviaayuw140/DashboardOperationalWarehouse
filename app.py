import os, sys, re, json, datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

BASE = Path(__file__).parent
EXCEL_GLOB = "*.xlsx"                  # semua file Excel di folder ini dicoba dibaca dan digabung (2025, 2026, dst.); file yang formatnya beda dilewati
DATA_JS = BASE / "assets" / "data.js"
SHEET_DATA = "Summary"                 # data transaksi gabungan inbound + outbound
SHEET_2025 = "Produktivitas 2025"      # ringkasan 2025 di file 2026 (dipakai hanya kalau tidak ada data transaksi 2025)
OFFICE_BASE_2025 = "LSB"
SHEET_PROD = "Produktivitas 2026"      # produktivitas per warehouse (CBM/MP dan transaksi/MP) + pembanding tahun lalu

COLS = {  # nama kolom standar -> kemungkinan nama di Excel (huruf besar/kecil tidak masalah)
    "Month": ["MONTH", "BULAN"], "Office": ["OFFICE", "WAREHOUSE", "GUDANG"],
    "TRANSACTION DATE": ["TRANSACTION DATE", "TRANSACTIONDATE", "TANGGAL TRANSAKSI", "TANGGAL"],
    "TRANSACTION": ["TRANSACTION", "TRANSACTION NO", "BOOKINGNUMBER", "BOOKING NUMBER"],
    "ACTIVITY": ["ACTIVITY", "AKTIVITAS"], "CLIENT": ["CLIENT", "CUSTOMER"],
    "VOLUME NEW": ["VOLUME NEW", "VOLUME-NEW", "VOLUMENEW", "VOLUME"], "QTY": ["QTY", "QUANTITY"],
}


def find_excels(args=None):
    if args:
        return [Path(a) for a in args]
    env = os.environ.get("WH_EXCEL")
    if env:
        return [Path(x) for x in env.split(os.pathsep) if x]
    return sorted(p for p in BASE.glob(EXCEL_GLOB) if not p.name.startswith("~$"))


# ---------- 1. Baca Excel ----------
def _rename(df):
    up = {str(c).strip().upper(): c for c in df.columns}
    out = {}
    for std, alts in COLS.items():
        for a in alts:
            if a in up:
                out[up[a]] = std
                break
    return df.rename(columns=out)


def _read_sheet(path):
    """Cari sheet data: pakai 'Summary' kalau ada, kalau tidak pakai sheet pertama yang kolomnya lengkap."""
    xl = pd.ExcelFile(path)
    names = [SHEET_DATA] + [n for n in xl.sheet_names if n != SHEET_DATA] if SHEET_DATA in xl.sheet_names else xl.sheet_names
    for n in names:
        head = _rename(xl.parse(n, nrows=3))
        if all(c in head.columns for c in COLS):
            return _rename(xl.parse(n)), n
    last = _rename(xl.parse(xl.sheet_names[0], nrows=3))
    raise ValueError(f"[{Path(path).name}] tidak ada sheet dengan kolom lengkap {list(COLS)}.\n"
                     f"Sheet: {xl.sheet_names}\nKolom di sheet pertama: {list(last.columns)}")


def load_lines(path):
    df, sheet = _read_sheet(path)

    def to_dt(v):  # sebagian tanggal tersimpan sebagai teks dd/mm/yyyy
        if isinstance(v, (dt.datetime, pd.Timestamp)):
            return pd.Timestamp(v)
        if isinstance(v, (int, float)) and not pd.isna(v):
            return pd.Timestamp("1899-12-30") + pd.to_timedelta(v, unit="D")
        return pd.to_datetime(str(v).strip(), dayfirst=True, errors="coerce")

    df["_d"] = df["TRANSACTION DATE"].map(to_dt)
    bad = int(df["_d"].isna().sum())
    df = df.dropna(subset=["_d"])
    df["act"] = df["ACTIVITY"].astype(str).str.lower().map(
        lambda a: "OUT" if re.search(r"out|deliver|dispatch|keluar", a) else ("IN" if re.search(r"in|receiv|masuk", a) else None))
    skipped = int(df["act"].isna().sum())
    df = df.dropna(subset=["act"])

    m = pd.to_numeric(df["Month"], errors="coerce")
    df["m"] = m.where(m.between(1, 12), df["_d"].dt.month).astype(int) - 1      # bulan mengikuti kolom Month
    df["y"] = df["_d"].dt.year
    df["trx"] = df["TRANSACTION"].fillna("(kosong)").astype(str).str.strip()
    df["cust"] = df["CLIENT"].fillna("(tanpa nama)").astype(str).str.strip()
    df["wh"] = df["Office"].fillna("-").astype(str).str.strip()
    df["vol"] = pd.to_numeric(df["VOLUME NEW"], errors="coerce").fillna(0)
    df["qty"] = pd.to_numeric(df["QTY"], errors="coerce").fillna(0)
    yrs = sorted(df["y"].unique().tolist())
    print(f"[{Path(path).name}] sheet '{sheet}': {len(df):,} baris item, tahun {yrs} (tanggal gagal {bad}, activity dilewati {skipped})")
    return df[["y", "m", "act", "cust", "wh", "trx", "vol", "qty"]]


def load_base_2025(path):
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return None
    if SHEET_2025 not in wb.sheetnames:
        return None
    vol = trx = None
    for r in wb[SHEET_2025].iter_rows(values_only=True):
        label = str(r[0] or "")
        vals = [v if isinstance(v, (int, float)) else None for v in r[2:14]]
        # ambil kemunculan PERTAMA saja (blok WH Surabaya); blok berikutnya di sheet yang sama adalah gudang lain
        if vol is None and re.match(r"throughput 2025", label, re.I):
            vol = vals
        elif trx is None and re.match(r"numb of line trans 2025", label, re.I):
            trx = vals
    return {"y": 2025, "wh": OFFICE_BASE_2025, "vol": vol, "trx": trx} if vol and trx else None


def load_productivity(path):
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return None
    if SHEET_PROD not in wb.sheetnames:
        return None
    year = int(re.search(r"\d{4}", SHEET_PROD).group())
    num = lambda v: float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None
    whs, cur, blk = [], None, None
    for r in wb[SHEET_PROD].iter_rows(values_only=True):
        lab = str(r[0] or "").strip()
        if lab.startswith("WH ") and isinstance(r[2], dt.datetime):
            name = lab[3:].strip()
            cur = next((w for w in whs if w["name"] == name), None)
            if cur is None:
                cur = {"name": name}
                whs.append(cur)
            blk = None
            continue
        if cur is None or not lab:
            continue
        cells = [num(v) for v in r[2:14]]
        if lab.startswith("Throughput"):
            blk = cur.setdefault("cbm", {}); blk["val"] = cells
        elif lab.startswith("Numb of Line Trans"):
            blk = cur.setdefault("trx", {}); blk["val"] = cells
        elif blk is not None and lab.startswith("Numb of MP"):
            blk["mp"] = cells
        elif blk is not None and (m := re.match(r"Productivity (\d{4})", lab)):
            if int(m.group(1)) == year:
                blk["prod"] = [c if (blk.get("val") and blk["val"][i] is not None) else None for i, c in enumerate(cells)]
            else:
                blk["prev"] = [c if (c and c > 0) else None for c in cells]
    out = []
    for w in whs:   # buang warehouse tanpa data sama sekali (mis. belum beroperasi)
        if any(v is not None for b in ("cbm", "trx") for v in w.get(b, {}).get("val", [])):
            for b in ("cbm", "trx"):
                for k in ("val", "mp", "prod", "prev"):
                    w.setdefault(b, {}).setdefault(k, [None] * 12)
            out.append(w)
    return {"year": year, "prev_year": year - 1, "wh": out} if out else None


def build_data(excels=None):
    files = [Path(f) for f in (excels or find_excels())]
    if not files:
        raise SystemExit(f"Tidak ada file Excel '{EXCEL_GLOB}' di {BASE.resolve()}")
    for f in files:
        if not f.exists():
            raise SystemExit(f"File Excel tidak ditemukan: {f.resolve()}")
    parts, used, errs = [], [], []
    for f in files:
        try:
            parts.append(load_lines(f))
            used.append(f)
        except ValueError as e:      # format tidak cocok (bukan file data transaksi): lewati
            errs.append(str(e))
            print("DILEWATI " + str(e).splitlines()[0])
    if not parts:
        raise SystemExit("Tidak ada file Excel yang formatnya cocok.\n" + "\n".join(errs))
    files = used
    df = pd.concat(parts, ignore_index=True)
    g = df.groupby(["y", "m", "act", "cust", "wh", "trx"], as_index=False).agg(vol=("vol", "sum"), qty=("qty", "sum"))
    years = sorted(g["y"].unique().tolist())
    base = None
    if len(years) < 2:     # tidak ada tahun pembanding di data transaksi -> pakai ringkasan 2025 kalau ada
        for f in files:
            base = load_base_2025(f)
            if base:
                break
    prod = next((p for p in (load_productivity(f) for f in files) if p), None)
    print(f"Digabung: {len(df):,} baris item -> {len(g):,} baris transaksi, tahun {years}"
          + (", pembanding 2025 dari sheet ringkasan (LSB)" if base else ""))
    rows = [{"y": int(r.y), "m": int(r.m), "act": r.act, "cust": r.cust, "wh": r.wh, "trx": r.trx,
             "vol": round(float(r.vol), 4), "qty": round(float(r.qty), 2)} for r in g.itertuples()]
    data = {"rows": rows, "base": base, "prod": prod, "source": " + ".join(f.name for f in files), "nlines": int(len(df)),
            "built": dt.datetime.now().strftime("%d/%m/%Y %H:%M")}
    DATA_JS.parent.mkdir(exist_ok=True)
    DATA_JS.write_text("window.DATA=" + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";", encoding="utf-8")
    print(f"Tulis {DATA_JS} ({DATA_JS.stat().st_size/1e6:.1f} MB). Buka index.html di browser.")


# ---------- 2. Streamlit: gabungkan index.html + css + js jadi satu halaman ----------
def inline_page():
    import base64, mimetypes
    html = (BASE / "index.html").read_text(encoding="utf-8")

    def css(m):
        p = BASE / m.group(1).split("?")[0]
        return "<style>" + p.read_text(encoding="utf-8") + "</style>" if p.exists() else m.group(0)

    def js(m):
        p = BASE / m.group(1).split("?")[0]
        if not p.exists():
            return m.group(0)
        return "<script>" + p.read_text(encoding="utf-8").replace("</script", "<\\/script") + "</script>"

    def img(m):  
        src = m.group(2)
        p = BASE / src.split("?")[0]
        if src.startswith("data:") or not p.exists():
            return m.group(0)
        mime = mimetypes.guess_type(p.name)[0] or "image/png"
        return m.group(1) + "data:" + mime + ";base64," + base64.b64encode(p.read_bytes()).decode() + m.group(3)

    html = re.sub(r'<link[^>]*href=["\']([^"\']+\.css[^"\']*)["\'][^>]*>', css, html)
    html = re.sub(r'<script[^>]*src=["\']([^"\']+)["\'][^>]*>\s*</script>', js, html)
    html = re.sub(r'(<img[^>]*\bsrc=["\'])([^"\']+)(["\'])', img, html)
    sisa = re.findall(r'<script[^>]*src=[^>]*>', html)
    if sisa:
        print("SCRIPT BELUM TERSISIP:", sisa)
    return html


def run_streamlit():
    import streamlit as st
    import streamlit.components.v1 as components
    st.set_page_config(page_title="Warehouse Dashboard", layout="wide")
    st.markdown("<style>.block-container{padding-top:1rem}</style>", unsafe_allow_html=True)
    files = find_excels()
    st.sidebar.write("File Excel yang dibaca:", [f.name for f in files] or "tidak ada")
    newest = max((f.stat().st_mtime for f in files if f.exists()), default=0)
    stale = not DATA_JS.exists() or DATA_JS.stat().st_mtime < newest
    if st.sidebar.button("Muat ulang dari Excel") or stale:
        with st.spinner("Membaca Excel…"):
            build_data(files)
    components.html(inline_page(), height=1900, scrolling=True)


if len(sys.argv) > 1 and sys.argv[1] == "build":
    build_data(sys.argv[2:] or None)
elif "streamlit" in sys.modules:      # dijalankan lewat: streamlit run app.py
    run_streamlit()
else:
    print(__doc__)